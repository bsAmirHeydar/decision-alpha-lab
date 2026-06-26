#!/usr/bin/env python3
"""
Build one Excel workbook for entry/exit windows across multiple astro execution families.

No Strategy Tester is required. This tool works directly from one deterministic astro CSV.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict
from pathlib import Path
from typing import Dict, List, Optional

try:
    from openpyxl import Workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required. Install it with: python -m pip install openpyxl") from exc

from astro_paper_family_runner import FAMILY_DEFAULTS, load_config  # type: ignore
from astro_pure_entry_excel import (  # type: ignore
    BAR_HEADERS,
    EXIT_HEADERS,
    SUMMARY_HEADERS,
    WINDOW_HEADERS,
    add_sheet,
    build_entry_windows,
    compute_records,
    iter_rows,
    resolve_report_thresholds,
    write_csv,
)


DEFAULT_FAMILIES = ["PURE", "A0001", "A0002", "A0003", "A0004", "A0005", "A0006", "A0090"]


def write_xlsx(path: Path, sheets: List[tuple[str, List[str], List[dict]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    if default:
        wb.remove(default)
    for name, headers, rows in sheets:
        add_sheet(wb, name, headers, rows)
    wb.save(path)


def family_summary_rows(family: str, windows: List[dict], bars: List[dict], exits: List[dict], thresholds: object) -> List[dict]:
    return [
        {"key": "family", "value": family},
        {"key": "family_name", "value": getattr(thresholds, "family_name", family)},
        {"key": "windows", "value": len(windows)},
        {"key": "entry_bars", "value": len(bars)},
        {"key": "exit_events", "value": len(exits)},
        {"key": "buy_windows", "value": sum(1 for w in windows if w.get("side") == "BUY")},
        {"key": "sell_windows", "value": sum(1 for w in windows if w.get("side") == "SELL")},
        {"key": "open_at_csv_end", "value": sum(1 for w in windows if int(w.get("open_at_csv_end", 0) or 0) == 1)},
        {"key": "thresholds_json", "value": json.dumps(asdict(thresholds), ensure_ascii=False)},
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description="Build one Excel workbook of entry/exit windows across astro families")
    parser.add_argument("--csv", required=True, help="Input astro feature CSV")
    parser.add_argument("--out-xlsx", required=True, help="Output Excel workbook path")
    parser.add_argument("--config", default="", help="Optional doctrine/threshold config JSON")
    parser.add_argument("--families", nargs="+", default=DEFAULT_FAMILIES, choices=["PURE", *sorted(FAMILY_DEFAULTS.keys())])
    parser.add_argument("--also-csv", action="store_true", help="Also write combined CSV exports next to workbook")
    args = parser.parse_args()

    csv_path = Path(args.csv)
    out_xlsx = Path(args.out_xlsx)
    if not csv_path.exists():
        raise SystemExit(f"Input CSV not found: {csv_path}")

    config: Optional[Dict[str, object]] = load_config(args.config) if args.config else None
    rows = list(iter_rows(csv_path))

    suite_summary: List[dict] = [
        {"key": "input_csv", "value": str(csv_path)},
        {"key": "output_xlsx", "value": str(out_xlsx)},
        {"key": "families", "value": ", ".join(args.families)},
        {"key": "rows_scanned", "value": len(rows)},
    ]
    family_overview: List[dict] = []
    all_windows: List[dict] = []
    all_bars: List[dict] = []
    all_exits: List[dict] = []
    workbook_sheets: List[tuple[str, List[str], List[dict]]] = []

    for family in args.families:
        thresholds = resolve_report_thresholds(family, config)
        records = compute_records(rows, family, thresholds)
        windows, bars, exits = build_entry_windows(records, rows, family, thresholds)

        for row in windows:
            tagged = dict(row)
            tagged["family"] = family
            all_windows.append(tagged)
        for row in bars:
            tagged = dict(row)
            tagged["family"] = family
            all_bars.append(tagged)
        for row in exits:
            tagged = dict(row)
            tagged["family"] = family
            all_exits.append(tagged)

        family_overview.append({
            "family": family,
            "family_name": thresholds.family_name,
            "windows": len(windows),
            "entry_bars": len(bars),
            "exit_events": len(exits),
            "buy_windows": sum(1 for w in windows if w.get("side") == "BUY"),
            "sell_windows": sum(1 for w in windows if w.get("side") == "SELL"),
            "open_at_csv_end": sum(1 for w in windows if int(w.get("open_at_csv_end", 0) or 0) == 1),
        })

        workbook_sheets.append((f"{family}_Summary", SUMMARY_HEADERS, family_summary_rows(family, windows, bars, exits, thresholds)))
        workbook_sheets.append((f"{family}_Windows", WINDOW_HEADERS, windows))
        workbook_sheets.append((f"{family}_Exits", EXIT_HEADERS, exits))

    overview_headers = ["family", "family_name", "windows", "entry_bars", "exit_events", "buy_windows", "sell_windows", "open_at_csv_end"]
    add_family_header_windows = WINDOW_HEADERS + (["family"] if "family" not in WINDOW_HEADERS else [])
    add_family_header_bars = ["family"] + BAR_HEADERS if "family" not in BAR_HEADERS else BAR_HEADERS
    add_family_header_exits = ["family"] + EXIT_HEADERS if "family" not in EXIT_HEADERS else EXIT_HEADERS

    sheets: List[tuple[str, List[str], List[dict]]] = [
        ("RunSummary", SUMMARY_HEADERS, suite_summary),
        ("FamilyOverview", overview_headers, family_overview),
        ("AllEntryWindows", add_family_header_windows, all_windows),
        ("AllEntryBars", add_family_header_bars, all_bars),
        ("AllExitEvents", add_family_header_exits, all_exits),
    ] + workbook_sheets

    write_xlsx(out_xlsx, sheets)

    if args.also_csv:
        stem_dir = out_xlsx.parent
        stem = out_xlsx.stem
        write_csv(stem_dir / f"{stem}_AllEntryWindows.csv", all_windows, add_family_header_windows)
        write_csv(stem_dir / f"{stem}_AllEntryBars.csv", all_bars, add_family_header_bars)
        write_csv(stem_dir / f"{stem}_AllExitEvents.csv", all_exits, add_family_header_exits)

    print(json.dumps({
        "ok": True,
        "input_csv": str(csv_path),
        "output_xlsx": str(out_xlsx),
        "families": args.families,
        "rows_scanned": len(rows),
        "all_entry_windows": len(all_windows),
        "all_entry_bars": len(all_bars),
        "all_exit_events": len(all_exits),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
