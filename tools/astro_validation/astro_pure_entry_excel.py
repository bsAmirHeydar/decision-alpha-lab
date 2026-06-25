#!/usr/bin/env python3
"""
Build an Excel report of complete Pure Astro entry windows.

Input:
  - raw astro feature CSV produced by tools/astro_feature_builder/astro_feature_builder.py
  - optionally a family gate profile: PURE, A0001, A0002, A0003, A0090

Output:
  - .xlsx workbook with:
      RunSummary    : metadata and counts
      EntryWindows  : one row per full enter_long / enter_short window
      EntryBars     : every bar that belongs to those full entry windows
      ExitEvents    : resolved exit/warning rows for each entry window

This is intentionally Python-side and batch-based: it does not use MT5 ticks.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import asdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - user environment message
    raise SystemExit(
        "openpyxl is required to write .xlsx. Install it with: python -m pip install openpyxl"
    ) from exc

# Reuse the same pure signal logic as the existing paper runner.
from astro_paper_family_runner import (  # type: ignore
    FAMILY_DEFAULTS,
    PureSignal,
    ThresholdProfile,
    ExecState,
    apply_family_gates,
    compute_signal,
    load_config,
    resolve_thresholds,
    safe,
    step_state,
)

FULL_ENTRIES = {"enter_long", "enter_short"}
OPPOSITE = {"enter_long": "enter_short", "enter_short": "enter_long"}


PURE_DEFAULTS = {
    "family_name": "PURE_pure_astro_signal",
    "arm_threshold": 58.0,
    "enter_threshold": 66.0,
    "reduce_threshold": 52.0,
    "exit_threshold": 58.0,
    "natal_activation_minimum": 0.0,
    "friction_minimum": 0.0,
}


SIGNAL_FIELDS = [
    "entry_score",
    "exit_score",
    "long_bias_score",
    "short_bias_score",
    "path_score",
    "friction_score",
    "volatility_score",
    "natal_activation_score",
    "macro_timing_score",
    "meso_timing_score",
    "micro_timing_score",
    "minute_window_score",
    "minute_exhaustion_score",
]


WINDOW_HEADERS = [
    "entry_id",
    "family",
    "side",
    "entry_signal",
    "direction",
    "regime",
    "state",
    "valid_from_broker",
    "valid_until_broker",
    "valid_until_exclusive_broker",
    "bars",
    "entry_start_utc",
    "entry_end_utc",
    "exit_time_broker",
    "exit_time_utc",
    "exit_after_bars",
    "exit_reason",
    "first_exit_warning_broker",
    "first_opposite_entry_broker",
    "open_at_csv_end",
    "entry_score_avg",
    "entry_score_min",
    "entry_score_max",
    "exit_score_at_exit",
    "path_score_avg",
    "friction_score_avg",
    "volatility_score_avg",
    "natal_activation_avg",
    "macro_timing_avg",
    "meso_timing_avg",
    "micro_timing_avg",
    "minute_window_avg",
    "minute_exhaustion_avg",
    "long_bias_avg",
    "short_bias_avg",
    "doctrine_id",
    "schema_version",
    "natal_label",
    "feature_key_start",
    "astro_trade_key_start",
    "astro_language_start",
    "window_key",
]


BAR_HEADERS = [
    "entry_id",
    "bar_index",
    "broker_time",
    "utc_time",
    "side",
    "entry_signal",
    "exit_signal",
    "direction",
    "regime",
    "state",
] + SIGNAL_FIELDS + [
    "macro_context",
    "meso_context",
    "micro_context",
    "minute_context",
    "feature_key",
    "astro_trade_key",
    "astro_language",
]


EXIT_HEADERS = [
    "entry_id",
    "event_type",
    "broker_time",
    "utc_time",
    "bar_index",
    "side",
    "entry_signal",
    "exit_signal",
    "direction",
    "regime",
    "state",
    "entry_score",
    "exit_score",
    "reason",
]


SUMMARY_HEADERS = ["key", "value"]


def iter_rows(path: Path) -> Iterable[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        yield from csv.DictReader(f)


def parse_time(text: str) -> Optional[datetime]:
    text = (text or "").strip()
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y.%m.%d %H:%M:%S",
        "%Y.%m.%d %H:%M",
        "%Y-%m-%dT%H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def fmt_time(value: Optional[datetime] | str) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def as_float_text(value: object) -> float:
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return 0.0


def avg(values: Sequence[float]) -> float:
    return sum(values) / len(values) if values else 0.0


def minimum(values: Sequence[float]) -> float:
    return min(values) if values else 0.0


def maximum(values: Sequence[float]) -> float:
    return max(values) if values else 0.0


def resolve_report_thresholds(family: str, config: Optional[Dict[str, object]]) -> ThresholdProfile:
    if family == "PURE":
        return ThresholdProfile(**PURE_DEFAULTS)
    return resolve_thresholds(family, config)


def maybe_apply_family(family: str, signal: PureSignal, row: dict, thresholds: ThresholdProfile) -> PureSignal:
    if family == "PURE":
        return signal
    return apply_family_gates(family, signal, row, thresholds)


def signal_key(signal: PureSignal) -> str:
    return (
        f"dir={signal.direction_name}|entry={signal.entry_signal}|exit={signal.exit_signal}"
        f"|regime={signal.regime_name}|state={signal.trigger_state}"
    )


def timeframe_seconds(rows: List[dict]) -> int:
    if len(rows) >= 2:
        t0 = parse_time(rows[0].get("broker_time", ""))
        t1 = parse_time(rows[1].get("broker_time", ""))
        if t0 and t1:
            delta = int((t1 - t0).total_seconds())
            if delta > 0:
                return delta
    return 60


def side_from_entry(entry_signal: str) -> str:
    if entry_signal == "enter_long":
        return "BUY"
    if entry_signal == "enter_short":
        return "SELL"
    return ""


def exit_reason_for(signal: PureSignal, start_entry: str, thresholds: ThresholdProfile) -> str:
    if signal.entry_signal == OPPOSITE.get(start_entry, ""):
        return "opposite_full_entry"
    if signal.trigger_state == "timing_exit":
        return "timing_exit"
    if signal.exit_signal == "exit_or_reduce" and signal.exit_score >= thresholds.exit_threshold:
        return "exit_score_threshold"
    if signal.exit_signal == "exit_or_reduce" and signal.exit_score >= thresholds.reduce_threshold:
        return "reduce_score_threshold"
    if signal.exit_signal == "exit_or_reduce":
        return "exit_or_reduce"
    return "state_machine_exit"


def compute_records(rows: List[dict], family: str, thresholds: ThresholdProfile) -> List[dict]:
    records: List[dict] = []
    state = ExecState(family_name=thresholds.family_name)
    for index, row in enumerate(rows):
        signal = compute_signal(row)
        signal = maybe_apply_family(family, signal, row, thresholds)
        state = step_state(state, signal, thresholds)
        records.append({
            "index": index,
            "row": row,
            "signal": signal,
            "state_phase": state.phase,
            "state_action": state.action,
            "state_position_direction": state.position_direction,
            "state_hold_bars": state.hold_bars,
        })
    return records


def find_exit_from(records: List[dict], start_index: int, start_entry: str, thresholds: ThresholdProfile) -> Tuple[Optional[int], str, Optional[int], Optional[int]]:
    """Return hard exit index, reason, first warning index, first opposite entry index."""
    first_warning: Optional[int] = None
    first_opposite: Optional[int] = None

    state = ExecState(family_name=thresholds.family_name)
    for i in range(start_index, len(records)):
        signal: PureSignal = records[i]["signal"]
        if i > start_index:
            if first_warning is None and signal.exit_signal == "exit_or_reduce":
                first_warning = i
            if first_opposite is None and signal.entry_signal == OPPOSITE.get(start_entry, ""):
                first_opposite = i

        state = step_state(state, signal, thresholds)
        if i > start_index and state.action == "exit":
            return i, exit_reason_for(signal, start_entry, thresholds), first_warning, first_opposite

    return None, "no_exit_inside_csv", first_warning, first_opposite


def build_entry_windows(records: List[dict], rows: List[dict], family: str, thresholds: ThresholdProfile) -> Tuple[List[dict], List[dict], List[dict]]:
    tf_sec = timeframe_seconds(rows)
    windows: List[dict] = []
    bars: List[dict] = []
    exits: List[dict] = []
    entry_id = 1
    i = 0
    while i < len(records):
        sig: PureSignal = records[i]["signal"]
        if sig.entry_signal not in FULL_ENTRIES:
            i += 1
            continue

        start = i
        entry_signal = sig.entry_signal
        key = signal_key(sig)
        j = i
        while j + 1 < len(records):
            nxt: PureSignal = records[j + 1]["signal"]
            if nxt.entry_signal != entry_signal or signal_key(nxt) != key:
                break
            j += 1

        win_records = records[start:j + 1]
        first_row = rows[start]
        last_row = rows[j]
        start_bt = parse_time(first_row.get("broker_time", ""))
        end_bt = parse_time(last_row.get("broker_time", ""))
        start_utc = parse_time(first_row.get("utc_time", ""))
        end_utc = parse_time(last_row.get("utc_time", ""))
        end_exclusive = (end_bt + timedelta(seconds=tf_sec)) if end_bt else ""

        exit_idx, exit_reason, warning_idx, opposite_idx = find_exit_from(records, start, entry_signal, thresholds)
        exit_record = records[exit_idx] if exit_idx is not None else None
        warning_record = records[warning_idx] if warning_idx is not None else None
        opposite_record = records[opposite_idx] if opposite_idx is not None else None
        exit_sig: Optional[PureSignal] = exit_record["signal"] if exit_record else None
        exit_row = rows[exit_idx] if exit_idx is not None else {}

        def values_for(field: str) -> List[float]:
            return [as_float_text(getattr(r["signal"], field)) for r in win_records]

        current_id = f"E{entry_id:05d}"
        summary = {
            "entry_id": current_id,
            "family": family,
            "side": side_from_entry(entry_signal),
            "entry_signal": entry_signal,
            "direction": sig.direction_name,
            "regime": sig.regime_name,
            "state": sig.trigger_state,
            "valid_from_broker": fmt_time(start_bt) or first_row.get("broker_time", ""),
            "valid_until_broker": fmt_time(end_bt) or last_row.get("broker_time", ""),
            "valid_until_exclusive_broker": fmt_time(end_exclusive),
            "bars": j - start + 1,
            "entry_start_utc": fmt_time(start_utc) or first_row.get("utc_time", ""),
            "entry_end_utc": fmt_time(end_utc) or last_row.get("utc_time", ""),
            "exit_time_broker": exit_row.get("broker_time", "") if exit_idx is not None else "",
            "exit_time_utc": exit_row.get("utc_time", "") if exit_idx is not None else "",
            "exit_after_bars": (exit_idx - start) if exit_idx is not None else "",
            "exit_reason": exit_reason,
            "first_exit_warning_broker": warning_record["row"].get("broker_time", "") if warning_record else "",
            "first_opposite_entry_broker": opposite_record["row"].get("broker_time", "") if opposite_record else "",
            "open_at_csv_end": 1 if exit_idx is None else 0,
            "entry_score_avg": avg(values_for("entry_score")),
            "entry_score_min": minimum(values_for("entry_score")),
            "entry_score_max": maximum(values_for("entry_score")),
            "exit_score_at_exit": exit_sig.exit_score if exit_sig else "",
            "path_score_avg": avg(values_for("path_score")),
            "friction_score_avg": avg(values_for("friction_score")),
            "volatility_score_avg": avg(values_for("volatility_score")),
            "natal_activation_avg": avg(values_for("natal_activation_score")),
            "macro_timing_avg": avg(values_for("macro_timing_score")),
            "meso_timing_avg": avg(values_for("meso_timing_score")),
            "micro_timing_avg": avg(values_for("micro_timing_score")),
            "minute_window_avg": avg(values_for("minute_window_score")),
            "minute_exhaustion_avg": avg(values_for("minute_exhaustion_score")),
            "long_bias_avg": avg(values_for("long_bias_score")),
            "short_bias_avg": avg(values_for("short_bias_score")),
            "doctrine_id": sig.doctrine_id,
            "schema_version": sig.schema_version,
            "natal_label": first_row.get("natal_label", ""),
            "feature_key_start": safe(first_row.get("feature_key", "")),
            "astro_trade_key_start": safe(sig.astro_trade_key),
            "astro_language_start": safe(sig.astro_language),
            "window_key": safe(key),
        }
        windows.append(summary)

        for r in win_records:
            s: PureSignal = r["signal"]
            source_row = r["row"]
            bar = {
                "entry_id": current_id,
                "bar_index": r["index"],
                "broker_time": source_row.get("broker_time", ""),
                "utc_time": source_row.get("utc_time", ""),
                "side": side_from_entry(s.entry_signal),
                "entry_signal": s.entry_signal,
                "exit_signal": s.exit_signal,
                "direction": s.direction_name,
                "regime": s.regime_name,
                "state": s.trigger_state,
                **{field: getattr(s, field) for field in SIGNAL_FIELDS},
                "macro_context": safe(s.macro_context),
                "meso_context": safe(s.meso_context),
                "micro_context": safe(s.micro_context),
                "minute_context": safe(s.minute_context),
                "feature_key": safe(source_row.get("feature_key", "")),
                "astro_trade_key": safe(s.astro_trade_key),
                "astro_language": safe(s.astro_language),
            }
            bars.append(bar)

        if warning_record:
            ws: PureSignal = warning_record["signal"]
            wr = warning_record["row"]
            exits.append({
                "entry_id": current_id,
                "event_type": "first_exit_warning",
                "broker_time": wr.get("broker_time", ""),
                "utc_time": wr.get("utc_time", ""),
                "bar_index": warning_record["index"],
                "side": side_from_entry(entry_signal),
                "entry_signal": ws.entry_signal,
                "exit_signal": ws.exit_signal,
                "direction": ws.direction_name,
                "regime": ws.regime_name,
                "state": ws.trigger_state,
                "entry_score": ws.entry_score,
                "exit_score": ws.exit_score,
                "reason": "first exit_or_reduce after entry start",
            })
        if exit_record:
            xs: PureSignal = exit_record["signal"]
            xr = exit_record["row"]
            exits.append({
                "entry_id": current_id,
                "event_type": "resolved_exit",
                "broker_time": xr.get("broker_time", ""),
                "utc_time": xr.get("utc_time", ""),
                "bar_index": exit_record["index"],
                "side": side_from_entry(entry_signal),
                "entry_signal": xs.entry_signal,
                "exit_signal": xs.exit_signal,
                "direction": xs.direction_name,
                "regime": xs.regime_name,
                "state": xs.trigger_state,
                "entry_score": xs.entry_score,
                "exit_score": xs.exit_score,
                "reason": exit_reason,
            })

        entry_id += 1
        i = j + 1

    return windows, bars, exits


def write_csv(path: Path, rows: List[dict], headers: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({h: row.get(h, "") for h in headers})


def add_sheet(wb: Workbook, name: str, headers: List[str], rows: List[dict]) -> None:
    ws = wb.create_sheet(title=name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D1D5DB")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for cell in ws[col_letter][1: min(ws.max_row, 200)]:
            value = cell.value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[col_letter].width = max(10, min(max_len + 2, 42))

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    # Color side and open-status columns for readability.
    side_col = headers.index("side") + 1 if "side" in headers else 0
    open_col = headers.index("open_at_csv_end") + 1 if "open_at_csv_end" in headers else 0
    for row_idx in range(2, ws.max_row + 1):
        if side_col:
            side_val = str(ws.cell(row=row_idx, column=side_col).value or "")
            if side_val == "BUY":
                ws.cell(row=row_idx, column=side_col).fill = PatternFill("solid", fgColor="DCFCE7")
            elif side_val == "SELL":
                ws.cell(row=row_idx, column=side_col).fill = PatternFill("solid", fgColor="FEE2E2")
        if open_col and str(ws.cell(row=row_idx, column=open_col).value) == "1":
            ws.cell(row=row_idx, column=open_col).fill = PatternFill("solid", fgColor="FEF3C7")


def write_xlsx(path: Path, summary_rows: List[dict], windows: List[dict], bars: List[dict], exits: List[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    if default:
        wb.remove(default)
    add_sheet(wb, "RunSummary", SUMMARY_HEADERS, summary_rows)
    add_sheet(wb, "EntryWindows", WINDOW_HEADERS, windows)
    add_sheet(wb, "EntryBars", BAR_HEADERS, bars)
    add_sheet(wb, "ExitEvents", EXIT_HEADERS, exits)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build Excel report for full Pure Astro buy/sell entry windows")
    parser.add_argument("--csv", required=True, help="Input astro feature CSV, e.g. Common\\Files\\astro_nas100_mql.csv")
    parser.add_argument("--out-xlsx", required=True, help="Output .xlsx path")
    parser.add_argument("--family", default="PURE", choices=["PURE", *sorted(FAMILY_DEFAULTS.keys())], help="PURE or executor gate family")
    parser.add_argument("--config", default="", help="Optional threshold config JSON")
    parser.add_argument("--also-csv", action="store_true", help="Also write EntryWindows/EntryBars/ExitEvents as CSV next to the .xlsx")
    args = parser.parse_args()

    csv_path = Path(args.csv)
    out_xlsx = Path(args.out_xlsx)
    if not csv_path.exists():
        raise SystemExit(f"Input CSV not found: {csv_path}")

    config = load_config(args.config) if args.config else None
    thresholds = resolve_report_thresholds(args.family, config)
    rows = list(iter_rows(csv_path))
    records = compute_records(rows, args.family, thresholds)
    windows, bars, exits = build_entry_windows(records, rows, args.family, thresholds)

    summary_rows = [
        {"key": "input_csv", "value": str(csv_path)},
        {"key": "output_xlsx", "value": str(out_xlsx)},
        {"key": "family", "value": args.family},
        {"key": "family_name", "value": thresholds.family_name},
        {"key": "rows_scanned", "value": len(rows)},
        {"key": "full_entry_windows", "value": len(windows)},
        {"key": "entry_bars", "value": len(bars)},
        {"key": "exit_events", "value": len(exits)},
        {"key": "buy_windows", "value": sum(1 for w in windows if w.get("side") == "BUY")},
        {"key": "sell_windows", "value": sum(1 for w in windows if w.get("side") == "SELL")},
        {"key": "open_at_csv_end", "value": sum(1 for w in windows if int(w.get("open_at_csv_end", 0) or 0) == 1)},
        {"key": "thresholds_json", "value": json.dumps(asdict(thresholds), ensure_ascii=False)},
    ]

    write_xlsx(out_xlsx, summary_rows, windows, bars, exits)

    if args.also_csv:
        stem_dir = out_xlsx.parent
        stem = out_xlsx.stem
        write_csv(stem_dir / f"{stem}_EntryWindows.csv", windows, WINDOW_HEADERS)
        write_csv(stem_dir / f"{stem}_EntryBars.csv", bars, BAR_HEADERS)
        write_csv(stem_dir / f"{stem}_ExitEvents.csv", exits, EXIT_HEADERS)

    print(json.dumps({
        "ok": True,
        "input_csv": str(csv_path),
        "output_xlsx": str(out_xlsx),
        "family": args.family,
        "rows_scanned": len(rows),
        "full_entry_windows": len(windows),
        "entry_bars": len(bars),
        "exit_events": len(exits),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
