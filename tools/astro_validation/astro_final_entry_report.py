#!/usr/bin/env python3
"""
Build a final astro-only entry report by aggregating votes from multiple execution families.

This is an offline decision surface:
- raw deterministic astro CSV in
- family gates applied per bar
- weighted consensus across families
- explicit purity vetoes
- final entry / exit windows out
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:
    from openpyxl import Workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required. Install it with: python -m pip install openpyxl") from exc

from astro_paper_family_runner import FAMILY_DEFAULTS, load_config  # type: ignore
from astro_pure_entry_excel import add_sheet, compute_records, iter_rows, resolve_report_thresholds, write_csv  # type: ignore


DEFAULT_FAMILIES = ["A0001", "A0002", "A0003", "A0004", "A0005", "A0006", "A0007"]
BASE_FAMILY_WEIGHTS = {
    "A0001": 1.00,
    "A0002": 1.00,
    "A0003": 0.95,
    "A0004": 1.10,
    "A0005": 1.05,
    "A0006": 1.10,
    "A0007": 1.10,
    "A0090": 1.20,
}
PURITY_PROFILES: Dict[str, Dict[str, float]] = {
    "pure_strict": {
        "min_families": 3,
        "min_weight": 3.15,
        "min_consensus_strength": 62.0,
        "min_entry_score": 66.0,
        "min_minute_window": 61.0,
        "max_minute_exhaustion": 55.0,
        "max_direction_conflict_weight": 1.20,
        "min_macro_timing": 60.0,
        "min_meso_timing": 56.0,
        "min_micro_timing": 56.0,
    },
    "pure_balanced": {
        "min_families": 3,
        "min_weight": 3.00,
        "min_consensus_strength": 56.0,
        "min_entry_score": 64.0,
        "min_minute_window": 58.0,
        "max_minute_exhaustion": 58.0,
        "max_direction_conflict_weight": 1.35,
        "min_macro_timing": 57.0,
        "min_meso_timing": 53.0,
        "min_micro_timing": 53.0,
    },
    "pure_probe": {
        "min_families": 2,
        "min_weight": 2.00,
        "min_consensus_strength": 48.0,
        "min_entry_score": 60.0,
        "min_minute_window": 55.0,
        "max_minute_exhaustion": 61.0,
        "max_direction_conflict_weight": 1.60,
        "min_macro_timing": 53.0,
        "min_meso_timing": 50.0,
        "min_micro_timing": 50.0,
    },
}

SUMMARY_HEADERS = ["key", "value"]
FAMILY_HEADERS = ["family", "family_name", "bars", "entry_votes", "weight"]
BAR_HEADERS = [
    "broker_time",
    "utc_time",
    "decision",
    "direction",
    "long_votes",
    "short_votes",
    "long_weight",
    "short_weight",
    "winning_weight",
    "agreement_ratio",
    "consensus_strength",
    "purity_state",
    "veto_reason",
    "entry_score_avg",
    "macro_timing_avg",
    "meso_timing_avg",
    "micro_timing_avg",
    "minute_window_avg",
    "minute_exhaustion_avg",
    "active_families",
    "supporting_families",
    "regime",
    "trigger_state",
    "astro_language_sample",
]
WINDOW_HEADERS = [
    "entry_id",
    "direction",
    "decision",
    "valid_from_broker",
    "valid_until_broker",
    "bars",
    "purity_state",
    "entry_score_avg",
    "macro_timing_avg",
    "meso_timing_avg",
    "micro_timing_avg",
    "minute_window_avg",
    "minute_exhaustion_avg",
    "winning_weight_avg",
    "agreement_ratio_avg",
    "consensus_strength_avg",
    "supporting_families_union",
    "exit_time_broker",
    "exit_reason",
]


def avg(values: List[float]) -> float:
    return sum(values) / len(values) if values else 0.0


def resolve_family_weights(profile_name: str) -> Dict[str, float]:
    weights = dict(BASE_FAMILY_WEIGHTS)
    if profile_name == "pure_strict":
        weights["A0004"] = 1.18
        weights["A0005"] = 1.12
        weights["A0006"] = 1.18
        weights["A0007"] = 1.18
        weights["A0003"] = 0.90
        weights["A0090"] = 1.00
    elif profile_name == "pure_probe":
        weights["A0001"] = 1.05
        weights["A0002"] = 1.05
        weights["A0003"] = 1.00
    return weights


def decide_bar(
    row: dict,
    family_records: Dict[str, List[dict]],
    index: int,
    profile: Dict[str, float],
    family_weights: Dict[str, float],
) -> dict:
    long_votes = 0
    short_votes = 0
    long_weight = 0.0
    short_weight = 0.0
    long_signals: List[dict] = []
    short_signals: List[dict] = []
    active_families: List[str] = []

    for family, records in family_records.items():
        record = records[index]
        signal = record["signal"]
        if signal.entry_signal == "enter_long":
            long_votes += 1
            long_weight += family_weights.get(family, 1.0)
            long_signals.append({"family": family, "signal": signal})
            active_families.append(family)
        elif signal.entry_signal == "enter_short":
            short_votes += 1
            short_weight += family_weights.get(family, 1.0)
            short_signals.append({"family": family, "signal": signal})
            active_families.append(family)

    total_weight = sum(family_weights.get(f, 1.0) for f in family_records.keys())
    direction = "flat"
    decision = "wait"
    winners: List[dict] = []
    winning_weight = 0.0
    losing_weight = 0.0
    if long_weight > short_weight:
        direction = "long"
        winners = long_signals
        winning_weight = long_weight
        losing_weight = short_weight
    elif short_weight > long_weight:
        direction = "short"
        winners = short_signals
        winning_weight = short_weight
        losing_weight = long_weight

    weighted_entry = [item["signal"].entry_score for item in winners]
    weighted_macro = [item["signal"].macro_timing_score for item in winners]
    weighted_meso = [item["signal"].meso_timing_score for item in winners]
    weighted_micro = [item["signal"].micro_timing_score for item in winners]
    weighted_minute = [item["signal"].minute_window_score for item in winners]
    weighted_exhaust = [item["signal"].minute_exhaustion_score for item in winners]

    entry_score_avg = avg(weighted_entry)
    macro_avg = avg(weighted_macro)
    meso_avg = avg(weighted_meso)
    micro_avg = avg(weighted_micro)
    minute_avg = avg(weighted_minute)
    exhaust_avg = avg(weighted_exhaust)

    agreement_ratio = (100.0 * winning_weight / total_weight) if total_weight > 0 else 0.0
    directional_total = long_weight + short_weight
    consensus_strength = (100.0 * abs(long_weight - short_weight) / directional_total) if directional_total > 0 else 0.0

    purity_state = "blocked"
    veto_reasons: List[str] = []
    winning_count = len(winners)

    if direction == "flat":
        veto_reasons.append("no_direction")
    else:
        if winning_count < int(profile["min_families"]):
            veto_reasons.append("low_family_count")
        if winning_weight < float(profile["min_weight"]):
            veto_reasons.append("low_weight")
        if consensus_strength < float(profile["min_consensus_strength"]):
            veto_reasons.append("weak_consensus")
        if entry_score_avg < float(profile["min_entry_score"]):
            veto_reasons.append("entry_score_low")
        if minute_avg < float(profile["min_minute_window"]):
            veto_reasons.append("minute_window_low")
        if exhaust_avg > float(profile["max_minute_exhaustion"]):
            veto_reasons.append("minute_exhaustion_high")
        if macro_avg < float(profile["min_macro_timing"]):
            veto_reasons.append("macro_timing_low")
        if meso_avg < float(profile["min_meso_timing"]):
            veto_reasons.append("meso_timing_low")
        if micro_avg < float(profile["min_micro_timing"]):
            veto_reasons.append("micro_timing_low")
        if losing_weight > float(profile["max_direction_conflict_weight"]):
            veto_reasons.append("direction_conflict")

        if not veto_reasons:
            decision = f"enter_{direction}"
            purity_state = "pure_entry"
        elif (
            winning_count >= max(2, int(profile["min_families"]) - 1)
            and winning_weight >= max(1.5, float(profile["min_weight"]) - 1.0)
            and entry_score_avg >= float(profile["min_entry_score"]) - 6.0
            and minute_avg >= float(profile["min_minute_window"]) - 4.0
            and macro_avg >= float(profile["min_macro_timing"]) - 4.0
        ):
            decision = f"armed_{direction}"
            purity_state = "probe"

    supporting = ",".join(item["family"] for item in winners)
    sample_signal = winners[0]["signal"] if winners else None
    return {
        "broker_time": row.get("broker_time", ""),
        "utc_time": row.get("utc_time", ""),
        "decision": decision,
        "direction": direction,
        "long_votes": long_votes,
        "short_votes": short_votes,
        "long_weight": round(long_weight, 4),
        "short_weight": round(short_weight, 4),
        "winning_weight": round(winning_weight, 4),
        "agreement_ratio": round(agreement_ratio, 4),
        "consensus_strength": round(consensus_strength, 4),
        "purity_state": purity_state,
        "veto_reason": ",".join(veto_reasons),
        "entry_score_avg": round(entry_score_avg, 4),
        "macro_timing_avg": round(macro_avg, 4),
        "meso_timing_avg": round(meso_avg, 4),
        "micro_timing_avg": round(micro_avg, 4),
        "minute_window_avg": round(minute_avg, 4),
        "minute_exhaustion_avg": round(exhaust_avg, 4),
        "active_families": ",".join(active_families),
        "supporting_families": supporting,
        "regime": sample_signal.regime_name if sample_signal else "mixed",
        "trigger_state": sample_signal.trigger_state if sample_signal else "standby",
        "astro_language_sample": sample_signal.astro_language if sample_signal else "",
    }


def build_final_windows(bar_rows: List[dict]) -> List[dict]:
    windows: List[dict] = []
    i = 0
    entry_id = 1
    while i < len(bar_rows):
        current = bar_rows[i]
        if current["decision"] not in {"enter_long", "enter_short"}:
            i += 1
            continue
        start = i
        decision = current["decision"]
        direction = current["direction"]
        j = i
        while j + 1 < len(bar_rows) and bar_rows[j + 1]["decision"] == decision:
            j += 1
        slice_rows = bar_rows[start : j + 1]
        support_union = sorted({name for row in slice_rows for name in row["supporting_families"].split(",") if name})
        exit_time = bar_rows[j + 1]["broker_time"] if j + 1 < len(bar_rows) else ""
        exit_reason = "consensus_fade" if j + 1 < len(bar_rows) else "csv_end"
        windows.append(
            {
                "entry_id": f"F{entry_id:05d}",
                "direction": direction,
                "decision": decision,
                "valid_from_broker": slice_rows[0]["broker_time"],
                "valid_until_broker": slice_rows[-1]["broker_time"],
                "bars": len(slice_rows),
                "purity_state": slice_rows[0]["purity_state"],
                "entry_score_avg": round(avg([float(r["entry_score_avg"]) for r in slice_rows]), 4),
                "macro_timing_avg": round(avg([float(r["macro_timing_avg"]) for r in slice_rows]), 4),
                "meso_timing_avg": round(avg([float(r["meso_timing_avg"]) for r in slice_rows]), 4),
                "micro_timing_avg": round(avg([float(r["micro_timing_avg"]) for r in slice_rows]), 4),
                "minute_window_avg": round(avg([float(r["minute_window_avg"]) for r in slice_rows]), 4),
                "minute_exhaustion_avg": round(avg([float(r["minute_exhaustion_avg"]) for r in slice_rows]), 4),
                "winning_weight_avg": round(avg([float(r["winning_weight"]) for r in slice_rows]), 4),
                "agreement_ratio_avg": round(avg([float(r["agreement_ratio"]) for r in slice_rows]), 4),
                "consensus_strength_avg": round(avg([float(r["consensus_strength"]) for r in slice_rows]), 4),
                "supporting_families_union": ",".join(support_union),
                "exit_time_broker": exit_time,
                "exit_reason": exit_reason,
            }
        )
        entry_id += 1
        i = j + 1
    return windows


def write_xlsx(path: Path, sheets: List[Tuple[str, List[str], List[dict]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    if default:
        wb.remove(default)
    for name, headers, rows in sheets:
        add_sheet(wb, name, headers, rows)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a final astro-only entry report from family consensus")
    parser.add_argument("--csv", required=True)
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--config", default="")
    parser.add_argument("--families", nargs="+", default=DEFAULT_FAMILIES, choices=sorted(FAMILY_DEFAULTS.keys()))
    parser.add_argument("--profile", choices=sorted(PURITY_PROFILES.keys()), default="pure_strict")
    parser.add_argument("--min-families", type=int, default=None)
    parser.add_argument("--min-weight", type=float, default=None)
    parser.add_argument("--min-consensus-strength", type=float, default=None)
    parser.add_argument("--min-entry-score", type=float, default=None)
    parser.add_argument("--min-minute-window", type=float, default=None)
    parser.add_argument("--max-minute-exhaustion", type=float, default=None)
    parser.add_argument("--also-csv", action="store_true")
    args = parser.parse_args()

    csv_path = Path(args.csv)
    out_xlsx = Path(args.out_xlsx)
    if not csv_path.exists():
        raise SystemExit(f"Input CSV not found: {csv_path}")

    profile = dict(PURITY_PROFILES[args.profile])
    if args.min_families is not None:
        profile["min_families"] = float(args.min_families)
    if args.min_weight is not None:
        profile["min_weight"] = args.min_weight
    if args.min_consensus_strength is not None:
        profile["min_consensus_strength"] = args.min_consensus_strength
    if args.min_entry_score is not None:
        profile["min_entry_score"] = args.min_entry_score
    if args.min_minute_window is not None:
        profile["min_minute_window"] = args.min_minute_window
    if args.max_minute_exhaustion is not None:
        profile["max_minute_exhaustion"] = args.max_minute_exhaustion

    family_weights = resolve_family_weights(args.profile)
    config: Optional[Dict[str, object]] = load_config(args.config) if args.config else None
    rows = list(iter_rows(csv_path))
    family_records: Dict[str, List[dict]] = {}
    family_summary: List[dict] = []
    for family in args.families:
        thresholds = resolve_report_thresholds(family, config)
        records = compute_records(rows, family, thresholds)
        family_records[family] = records
        family_summary.append(
            {
                "family": family,
                "family_name": thresholds.family_name,
                "bars": len(records),
                "entry_votes": sum(1 for r in records if r["signal"].entry_signal in {"enter_long", "enter_short"}),
                "weight": family_weights.get(family, 1.0),
            }
        )

    bar_rows = [decide_bar(row, family_records, index, profile, family_weights) for index, row in enumerate(rows)]
    final_windows = build_final_windows(bar_rows)
    summary_rows = [
        {"key": "input_csv", "value": str(csv_path)},
        {"key": "output_xlsx", "value": str(out_xlsx)},
        {"key": "families", "value": ", ".join(args.families)},
        {"key": "profile", "value": args.profile},
        {"key": "profile_json", "value": json.dumps(profile, ensure_ascii=False)},
        {"key": "rows_scanned", "value": len(rows)},
        {"key": "final_entry_windows", "value": len(final_windows)},
        {"key": "enter_long_bars", "value": sum(1 for row in bar_rows if row["decision"] == "enter_long")},
        {"key": "enter_short_bars", "value": sum(1 for row in bar_rows if row["decision"] == "enter_short")},
        {"key": "armed_long_bars", "value": sum(1 for row in bar_rows if row["decision"] == "armed_long")},
        {"key": "armed_short_bars", "value": sum(1 for row in bar_rows if row["decision"] == "armed_short")},
        {"key": "pure_entry_bars", "value": sum(1 for row in bar_rows if row["purity_state"] == "pure_entry")},
        {"key": "blocked_bars", "value": sum(1 for row in bar_rows if row["purity_state"] == "blocked")},
    ]

    sheets = [
        ("FinalSummary", SUMMARY_HEADERS, summary_rows),
        ("FamilyWeights", FAMILY_HEADERS, family_summary),
        ("BarConsensus", BAR_HEADERS, bar_rows),
        ("FinalEntryWindows", WINDOW_HEADERS, final_windows),
    ]
    write_xlsx(out_xlsx, sheets)

    if args.also_csv:
        write_csv(out_xlsx.with_name(out_xlsx.stem + "_BarConsensus.csv"), bar_rows, BAR_HEADERS)
        write_csv(out_xlsx.with_name(out_xlsx.stem + "_FinalEntryWindows.csv"), final_windows, WINDOW_HEADERS)

    print(
        json.dumps(
            {
                "ok": True,
                "input_csv": str(csv_path),
                "output_xlsx": str(out_xlsx),
                "families": args.families,
                "profile": args.profile,
                "rows_scanned": len(rows),
                "final_entry_windows": len(final_windows),
                "enter_long_bars": sum(1 for row in bar_rows if row["decision"] == "enter_long"),
                "enter_short_bars": sum(1 for row in bar_rows if row["decision"] == "enter_short"),
                "pure_entry_bars": sum(1 for row in bar_rows if row["purity_state"] == "pure_entry"),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
